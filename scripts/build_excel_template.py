#!/usr/bin/env python3
"""Genera la plantilla Excel CajaClara para control de finanzas personales."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "CajaClara_Finanzas.xlsx"

BRAND = "1B5E4A"
INK = "1A2421"
MUTED = "5A6B64"
LINE = "D5DED9"
SURFACE = "F3F7F5"
WHITE = "FFFFFF"
INCOME = "2A7A5C"
EXPENSE = "B54A3A"
WARN = "C47A3A"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(size=11, bold=False, color=INK, name="Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, color=color)


def style_header_row(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill(BRAND)
        cell.font = font(11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def money_format(cell) -> None:
    cell.number_format = '"$"#,##0.00'
    cell.alignment = Alignment(horizontal="right")


def pct_format(cell) -> None:
    cell.number_format = "0%"
    cell.alignment = Alignment(horizontal="center")


def build() -> None:
    wb = Workbook()

    categories = [
        ("Salario", "Ingreso", "2A7A5C"),
        ("Freelance", "Ingreso", "3D8B6E"),
        ("Inversiones", "Ingreso", "5BA88A"),
        ("Vivienda", "Gasto", "4A6B8A"),
        ("Alimentación", "Gasto", "C47A3A"),
        ("Transporte", "Gasto", "6B7C8A"),
        ("Servicios", "Gasto", "7A6B9A"),
        ("Salud", "Gasto", "B54A5A"),
        ("Ocio", "Gasto", "3A8A9A"),
        ("Compras", "Gasto", "9A6B4A"),
        ("Otros", "Ambos", "6B726E"),
    ]

    accounts = [
        ("Cuenta corriente", "Corriente", 2840.50),
        ("Ahorros", "Ahorro", 12500.00),
        ("Efectivo", "Efectivo", 185.25),
    ]

    budgets = [
        ("Vivienda", 650),
        ("Alimentación", 400),
        ("Transporte", 120),
        ("Servicios", 150),
        ("Ocio", 100),
        ("Compras", 200),
    ]

    transactions = [
        ("2026-10-01", "Nómina octubre", "Ingreso", "Salario", "Cuenta corriente", 2450.00),
        ("2026-10-02", "Alquiler", "Gasto", "Vivienda", "Cuenta corriente", 620.00),
        ("2026-10-04", "Supermercado", "Gasto", "Alimentación", "Cuenta corriente", 87.40),
        ("2026-10-06", "Proyecto web", "Ingreso", "Freelance", "Cuenta corriente", 480.00),
        ("2026-10-07", "Gasolina", "Gasto", "Transporte", "Efectivo", 42.50),
        ("2026-10-08", "Luz y agua", "Gasto", "Servicios", "Cuenta corriente", 98.20),
        ("2026-10-10", "Cena fuera", "Gasto", "Alimentación", "Cuenta corriente", 54.90),
        ("2026-10-12", "Dividendos ETF", "Ingreso", "Inversiones", "Ahorros", 65.30),
        ("2026-10-14", "Farmacia", "Gasto", "Salud", "Efectivo", 28.75),
        ("2026-10-15", "Streaming + cine", "Gasto", "Ocio", "Cuenta corriente", 36.99),
        ("2026-10-18", "Ropa otoño", "Gasto", "Compras", "Cuenta corriente", 119.90),
        ("2026-10-20", "Mercado semanal", "Gasto", "Alimentación", "Cuenta corriente", 62.15),
        ("2026-10-21", "Taxi aeropuerto", "Gasto", "Transporte", "Efectivo", 28.00),
        ("2026-10-22", "Consultoría 2h", "Ingreso", "Freelance", "Cuenta corriente", 160.00),
        ("2026-10-23", "Internet fibra", "Gasto", "Servicios", "Cuenta corriente", 39.90),
    ]

    monthly = [
        ("May", 2680, 1920),
        ("Jun", 2750, 2105),
        ("Jul", 2900, 1980),
        ("Ago", 2550, 2240),
        ("Sep", 3100, 2050),
        ("Oct", 3155.30, 1218.69),
    ]

    # ——— Instrucciones ———
    ws_help = wb.active
    ws_help.title = "Instrucciones"
    ws_help["A1"] = "CajaClara"
    ws_help["A1"].font = font(28, bold=True, color=BRAND)
    ws_help["A2"] = "Plantilla Excel para control de finanzas personales"
    ws_help["A2"].font = font(14, color=MUTED)

    lines = [
        (4, "Cómo usarla", True),
        (5, "1. Abre Movimientos y registra cada ingreso o gasto (fecha, descripción, tipo, categoría, cuenta y monto).", False),
        (6, "2. Ajusta Categorías y Cuentas. Puedes añadir filas al final de cada lista.", False),
        (7, "3. Define límites mensuales en Presupuestos.", False),
        (8, "4. En Dashboard cambia Mes (B2) y Año (D2): los KPIs, el pastel y los presupuestos se recalculan solos.", False),
        (9, "5. Edita Histórico_mensual para el gráfico de barras de flujo de caja.", False),
        (11, "Hojas", True),
        (12, "• Dashboard — resumen del mes con KPIs y gráficos", False),
        (13, "• Movimientos — registro de transacciones (fuente de verdad)", False),
        (14, "• Categorías / Cuentas / Presupuestos — catálogos y límites", False),
        (15, "• Histórico_mensual — ingresos y gastos por mes", False),
    ]
    for row, text, is_title in lines:
        cell = ws_help.cell(row=row, column=1, value=text)
        cell.font = font(14, bold=True, color=BRAND) if is_title else font(11, color=INK)

    set_widths(ws_help, {"A": 110})
    ws_help.row_dimensions[1].height = 36

    # ——— Categorías ———
    ws_cat = wb.create_sheet("Categorías")
    ws_cat["A1"] = "Categoría"
    ws_cat["B1"] = "Tipo"
    ws_cat["C1"] = "Color (hex)"
    style_header_row(ws_cat, 1, 3)
    for i, (name, tipo, color) in enumerate(categories, start=2):
        ws_cat.cell(row=i, column=1, value=name).border = thin
        ws_cat.cell(row=i, column=2, value=tipo).border = thin
        c = ws_cat.cell(row=i, column=3, value=color)
        c.border = thin
        c.fill = fill(color)
        c.font = font(10, color=WHITE)
    last_cat = 1 + len(categories)
    set_widths(ws_cat, {"A": 18, "B": 12, "C": 14})

    dv_tipo = DataValidation(type="list", formula1='"Ingreso,Gasto,Ambos"', allow_blank=True)
    ws_cat.add_data_validation(dv_tipo)
    dv_tipo.add(f"B2:B{last_cat + 50}")

    # ——— Cuentas ———
    ws_acc = wb.create_sheet("Cuentas")
    ws_acc["A1"] = "Cuenta"
    ws_acc["B1"] = "Tipo"
    ws_acc["C1"] = "Saldo"
    style_header_row(ws_acc, 1, 3)
    for i, (name, tipo, bal) in enumerate(accounts, start=2):
        ws_acc.cell(row=i, column=1, value=name).border = thin
        ws_acc.cell(row=i, column=2, value=tipo).border = thin
        c = ws_acc.cell(row=i, column=3, value=bal)
        c.border = thin
        money_format(c)
    last_acc = 1 + len(accounts)
    total_row = last_acc + 1
    ws_acc.cell(row=total_row, column=1, value="TOTAL").font = font(11, bold=True, color=BRAND)
    total_acc = ws_acc.cell(row=total_row, column=3, value=f"=SUM(C2:C{last_acc})")
    total_acc.font = font(12, bold=True, color=BRAND)
    money_format(total_acc)
    set_widths(ws_acc, {"A": 20, "B": 12, "C": 14})

    dv_acc_tipo = DataValidation(type="list", formula1='"Corriente,Ahorro,Efectivo"', allow_blank=True)
    ws_acc.add_data_validation(dv_acc_tipo)
    dv_acc_tipo.add(f"B2:B{last_acc + 50}")

    # ——— Presupuestos ———
    ws_bud = wb.create_sheet("Presupuestos")
    for col, h in enumerate(["Categoría", "Límite mensual", "Gastado (mes)", "% usado", "Estado"], start=1):
        ws_bud.cell(row=1, column=col, value=h)
    style_header_row(ws_bud, 1, 5)

    for i, (cat, limit) in enumerate(budgets, start=2):
        ws_bud.cell(row=i, column=1, value=cat).border = thin
        lim = ws_bud.cell(row=i, column=2, value=limit)
        lim.border = thin
        money_format(lim)
        spent = ws_bud.cell(
            row=i,
            column=3,
            value=(
                f'=SUMIFS(Movimientos!$F:$F,Movimientos!$D:$D,A{i},Movimientos!$C:$C,"Gasto",'
                f'Movimientos!$A:$A,">="&DATE(Dashboard!$B$3,Dashboard!$B$2,1),'
                f'Movimientos!$A:$A,"<"&DATE(Dashboard!$B$3,Dashboard!$B$2+1,1))'
            ),
        )
        spent.border = thin
        money_format(spent)
        pct = ws_bud.cell(row=i, column=4, value=f"=IF(B{i}=0,0,C{i}/B{i})")
        pct.border = thin
        pct_format(pct)
        estado = ws_bud.cell(
            row=i,
            column=5,
            value=f'=IF(D{i}>1,"Excedido",IF(D{i}>=0.85,"Alerta","OK"))',
        )
        estado.border = thin
        estado.alignment = Alignment(horizontal="center")

    last_bud = 1 + len(budgets)
    ws_bud.conditional_formatting.add(
        f"E2:E{last_bud}",
        FormulaRule(formula=['E2="Excedido"'], fill=fill("F8D7D3"), font=font(11, bold=True, color=EXPENSE)),
    )
    ws_bud.conditional_formatting.add(
        f"E2:E{last_bud}",
        FormulaRule(formula=['E2="Alerta"'], fill=fill("FBE8D5"), font=font(11, bold=True, color=WARN)),
    )
    ws_bud.conditional_formatting.add(
        f"E2:E{last_bud}",
        FormulaRule(formula=['E2="OK"'], fill=fill("D8EDE4"), font=font(11, bold=True, color=INCOME)),
    )
    set_widths(ws_bud, {"A": 16, "B": 16, "C": 16, "D": 12, "E": 12})

    # ——— Histórico ———
    ws_hist = wb.create_sheet("Histórico_mensual")
    for col, h in enumerate(["Mes", "Ingresos", "Gastos", "Balance"], start=1):
        ws_hist.cell(row=1, column=col, value=h)
    style_header_row(ws_hist, 1, 4)
    for i, (month, inc, exp) in enumerate(monthly, start=2):
        ws_hist.cell(row=i, column=1, value=month).border = thin
        for col, val in ((2, inc), (3, exp)):
            c = ws_hist.cell(row=i, column=col, value=val)
            c.border = thin
            money_format(c)
        bal = ws_hist.cell(row=i, column=4, value=f"=B{i}-C{i}")
        bal.border = thin
        money_format(bal)
    last_hist = 1 + len(monthly)
    set_widths(ws_hist, {"A": 10, "B": 14, "C": 14, "D": 14})

    # ——— Movimientos ———
    ws_tx = wb.create_sheet("Movimientos")
    for col, h in enumerate(["Fecha", "Descripción", "Tipo", "Categoría", "Cuenta", "Monto"], start=1):
        ws_tx.cell(row=1, column=col, value=h)
    style_header_row(ws_tx, 1, 6)

    for i, (fecha, desc, tipo, cat, cuenta, monto) in enumerate(transactions, start=2):
        d = ws_tx.cell(row=i, column=1, value=fecha)
        d.number_format = "YYYY-MM-DD"
        d.border = thin
        ws_tx.cell(row=i, column=2, value=desc).border = thin
        ws_tx.cell(row=i, column=3, value=tipo).border = thin
        ws_tx.cell(row=i, column=4, value=cat).border = thin
        ws_tx.cell(row=i, column=5, value=cuenta).border = thin
        m = ws_tx.cell(row=i, column=6, value=monto)
        m.border = thin
        money_format(m)
        m.font = font(11, color=INCOME if tipo == "Ingreso" else EXPENSE)

    blank_start = 2 + len(transactions)
    blank_end = blank_start + 199
    for i in range(blank_start, blank_end + 1):
        for col in range(1, 7):
            ws_tx.cell(row=i, column=col).border = thin
        ws_tx.cell(row=i, column=1).number_format = "YYYY-MM-DD"
        money_format(ws_tx.cell(row=i, column=6))

    dv_tx_tipo = DataValidation(type="list", formula1='"Ingreso,Gasto"', allow_blank=True)
    dv_tx_cat = DataValidation(type="list", formula1=f"=Categorías!$A$2:$A${last_cat}", allow_blank=True)
    dv_tx_acc = DataValidation(type="list", formula1=f"=Cuentas!$A$2:$A${last_acc}", allow_blank=True)
    for dv, rng in (
        (dv_tx_tipo, f"C2:C{blank_end}"),
        (dv_tx_cat, f"D2:D{blank_end}"),
        (dv_tx_acc, f"E2:E{blank_end}"),
    ):
        ws_tx.add_data_validation(dv)
        dv.add(rng)

    set_widths(ws_tx, {"A": 12, "B": 28, "C": 10, "D": 14, "E": 18, "F": 12})
    ws_tx.auto_filter.ref = f"A1:F{blank_end}"
    ws_tx.freeze_panes = "A2"

    # ——— Dashboard ———
    ws_dash = wb.create_sheet("Dashboard")

    ws_dash.merge_cells("A1:H1")
    ws_dash["A1"] = "CajaClara · Dashboard financiero"
    ws_dash["A1"].font = font(22, bold=True, color=WHITE)
    ws_dash["A1"].fill = fill(BRAND)
    ws_dash["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws_dash.row_dimensions[1].height = 42

    ws_dash["A2"] = "Mes"
    ws_dash["B2"] = 10
    ws_dash["C2"] = "Año"
    ws_dash["D2"] = 2026
    ws_dash["A3"] = "Año (fórmulas)"
    ws_dash["B3"] = "=D2"
    for cell in ("A2", "C2", "A3"):
        ws_dash[cell].font = font(10, bold=True, color=MUTED)
    for cell in ("B2", "D2", "B3"):
        ws_dash[cell].font = font(12, bold=True, color=BRAND)
        ws_dash[cell].fill = fill(SURFACE)
        ws_dash[cell].border = thin
        ws_dash[cell].alignment = Alignment(horizontal="center")
    ws_dash["E2"] = "Cambia Mes (B2) y Año (D2) para filtrar el resumen."
    ws_dash["E2"].font = font(10, color=MUTED)

    # KPIs
    ws_dash["A5"] = "Balance total (cuentas)"
    ws_dash["C5"] = "Ingresos del mes"
    ws_dash["E5"] = "Gastos del mes"
    ws_dash["G5"] = "Tasa de ahorro"
    for addr in ("A5", "C5", "E5", "G5"):
        ws_dash[addr].font = font(10, color=MUTED)

    ws_dash.merge_cells("A5:B5")
    ws_dash.merge_cells("C5:D5")
    ws_dash.merge_cells("E5:F5")
    ws_dash.merge_cells("G5:H5")
    ws_dash.merge_cells("A6:B6")
    ws_dash.merge_cells("C6:D6")
    ws_dash.merge_cells("E6:F6")
    ws_dash.merge_cells("G6:H6")

    ws_dash["A6"] = f"=Cuentas!C{total_row}"
    ws_dash["C6"] = (
        '=SUMIFS(Movimientos!F:F,Movimientos!C:C,"Ingreso",'
        'Movimientos!A:A,">="&DATE(B3,B2,1),'
        'Movimientos!A:A,"<"&DATE(B3,B2+1,1))'
    )
    ws_dash["E6"] = (
        '=SUMIFS(Movimientos!F:F,Movimientos!C:C,"Gasto",'
        'Movimientos!A:A,">="&DATE(B3,B2,1),'
        'Movimientos!A:A,"<"&DATE(B3,B2+1,1))'
    )
    ws_dash["G6"] = "=IF(C6=0,0,(C6-E6)/C6)"

    for addr in ("A6", "C6", "E6"):
        money_format(ws_dash[addr])
        ws_dash[addr].fill = fill(WHITE)
        ws_dash[addr].border = thin
    ws_dash["A6"].font = font(18, bold=True, color=INK)
    ws_dash["C6"].font = font(18, bold=True, color=INCOME)
    ws_dash["E6"].font = font(18, bold=True, color=EXPENSE)
    pct_format(ws_dash["G6"])
    ws_dash["G6"].font = font(18, bold=True, color=BRAND)
    ws_dash["G6"].fill = fill(WHITE)
    ws_dash["G6"].border = thin
    ws_dash.row_dimensions[6].height = 32

    # Gastos por categoría
    ws_dash["A8"] = "Gastos por categoría (mes)"
    ws_dash["A8"].font = font(14, bold=True, color=BRAND)
    ws_dash.merge_cells("A8:B8")
    ws_dash["A9"] = "Categoría"
    ws_dash["B9"] = "Gastado"
    style_header_row(ws_dash, 9, 2)

    expense_cats = [c[0] for c in categories if c[1] in ("Gasto", "Ambos")]
    for i, cat in enumerate(expense_cats, start=10):
        ws_dash.cell(row=i, column=1, value=cat).border = thin
        cell = ws_dash.cell(
            row=i,
            column=2,
            value=(
                f'=SUMIFS(Movimientos!$F:$F,Movimientos!$D:$D,A{i},Movimientos!$C:$C,"Gasto",'
                f'Movimientos!$A:$A,">="&DATE($B$3,$B$2,1),'
                f'Movimientos!$A:$A,"<"&DATE($B$3,$B$2+1,1))'
            ),
        )
        cell.border = thin
        money_format(cell)
    last_exp_row = 9 + len(expense_cats)

    pie = PieChart()
    pie.title = "Distribución de gastos"
    labels = Reference(ws_dash, min_col=1, min_row=10, max_row=last_exp_row)
    data = Reference(ws_dash, min_col=2, min_row=9, max_row=last_exp_row)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = False
    pie.width = 12
    pie.height = 8
    ws_dash.add_chart(pie, "D8")

    # Presupuestos en dashboard
    bud_start = last_exp_row + 2
    ws_dash.cell(row=bud_start, column=1, value="Presupuestos del mes").font = font(14, bold=True, color=BRAND)
    ws_dash.merge_cells(f"A{bud_start}:E{bud_start}")
    for col, h in enumerate(["Categoría", "Límite", "Gastado", "% usado", "Estado"], start=1):
        cell = ws_dash.cell(row=bud_start + 1, column=col, value=h)
        cell.fill = fill(BRAND)
        cell.font = font(11, bold=True, color=WHITE)
        cell.border = thin
        cell.alignment = Alignment(horizontal="center")

    for i in range(len(budgets)):
        r = bud_start + 2 + i
        src = 2 + i
        for col in range(1, 6):
            cell = ws_dash.cell(row=r, column=col, value=f"=Presupuestos!{get_column_letter(col)}{src}")
            cell.border = thin
        money_format(ws_dash.cell(row=r, column=2))
        money_format(ws_dash.cell(row=r, column=3))
        pct_format(ws_dash.cell(row=r, column=4))
        ws_dash.cell(row=r, column=5).alignment = Alignment(horizontal="center")

    bud_end = bud_start + 1 + len(budgets)
    ws_dash.conditional_formatting.add(
        f"E{bud_start + 2}:E{bud_end}",
        FormulaRule(
            formula=[f'E{bud_start + 2}="Excedido"'],
            fill=fill("F8D7D3"),
            font=font(11, bold=True, color=EXPENSE),
        ),
    )
    ws_dash.conditional_formatting.add(
        f"E{bud_start + 2}:E{bud_end}",
        FormulaRule(
            formula=[f'E{bud_start + 2}="Alerta"'],
            fill=fill("FBE8D5"),
            font=font(11, bold=True, color=WARN),
        ),
    )
    ws_dash.conditional_formatting.add(
        f"E{bud_start + 2}:E{bud_end}",
        FormulaRule(
            formula=[f'E{bud_start + 2}="OK"'],
            fill=fill("D8EDE4"),
            font=font(11, bold=True, color=INCOME),
        ),
    )

    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Flujo de caja (histórico)"
    data_ref = Reference(ws_hist, min_col=2, max_col=3, min_row=1, max_row=last_hist)
    cats_ref = Reference(ws_hist, min_col=1, min_row=2, max_row=last_hist)
    bar.add_data(data_ref, titles_from_data=True)
    bar.set_categories(cats_ref)
    bar.shape = 4
    bar.width = 15
    bar.height = 9
    ws_dash.add_chart(bar, "D20")

    note_row = bud_end + 2
    ws_dash.cell(row=note_row, column=1, value="Cómo registrar movimientos").font = font(14, bold=True, color=BRAND)
    ws_dash.merge_cells(f"A{note_row}:F{note_row}")
    ws_dash.cell(
        row=note_row + 1,
        column=1,
        value="Usa la hoja Movimientos (listas desplegables en Tipo, Categoría y Cuenta). El Dashboard se actualiza con fórmulas.",
    ).font = font(10, color=MUTED)

    set_widths(ws_dash, {"A": 18, "B": 14, "C": 14, "D": 14, "E": 14, "F": 12, "G": 12, "H": 12})

    order = [
        "Instrucciones",
        "Dashboard",
        "Movimientos",
        "Categorías",
        "Cuentas",
        "Presupuestos",
        "Histórico_mensual",
    ]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    wb.save(OUT)
    print(f"Escrito: {OUT}")


if __name__ == "__main__":
    build()
